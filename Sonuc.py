import openpyxl
import os
from openpyxl.styles import Alignment
from openpyxl.styles import Font
import datetime

HEADERS = {
    "A": ("Device SN\\Test Adımları", 30),
    "B": ("Tarih - Zaman", 30),
    "C": ("Data Select Eksik Veriler", 30),
    "D": ("Data Select Fazla Veriler", 30),
    "E": ("Data Select Testi Sonucu", 30),
    "F": ("SN Eşleşmesi (Limit Değerleri ve Sistem)", 45),
    "G": ("FW Eşleşmesi (Limit Değerleri ve Sistem)", 45),
    "H": ("DN Eşleşmesi (Limit Değerleri ve Sistem)", 45),
    "I": ("HW Eşleşmesi (Limit Değerleri ve Sistem)", 45),
    "J": ("Acc Kalib Yüklenmesi", 30),
    "K": ("Gyro Kalib Yüklenmesi", 30),
    "L": ("Mag Kalib Yüklenmesi", 30),
    "M": ("Kalibrasyon - Config Kontrol Sonucu", 45),
    "N": ("Hard Reset - 1 Sonucu:", 30),
    "O": ("Hard Reset - 2 Sonucu:", 30),
    "P": ("Soft Reset - 1 Sonucu:", 30),
    "Q": ("Soft Reset - 2 Sonucu:", 30),
    "R": ("Reset Testi Sonucu:", 30),
    "S": ("Acc Norm Açılış Sonucu:", 45),
    "T": ("Max Acc Norm Değeri:", 45),
    "U": ("Min Acc Norm Değeri:", 45),
    "V": ("Acc Norm Status:", 50),
    "W": ("Gyro Açılış Sonucu:", 45),
    "X": ("Acc Döndürme Test - 1 Sonucu:", 40),
    "Y": ("Acc Döndürme Test - 2 Sonucu:", 40),
    "Z": ("Acc Döndürme Test - 3 Sonucu:", 40),
    "AA": ("Acc Döndürme Test - 4 Sonucu:", 40),
    "AB": ("Acc Döndürme Test - 5 Sonucu:", 40),
    "AC": ("Acc Döndürme Test - 6 Sonucu:", 40),
    "AD": ("Acc Döndürme Test Sonucu:", 40),
    "AE": ("Konum 1 Roll Ortalaması:", 40),
    "AF": ("Konum 1 Pitch Ortalaması:", 40),
    "AG": ("Konum 2 Roll Ortalaması:", 40),
    "AH": ("Konum 2 Pitch Ortalaması:", 40),
    "AI": ("Konum 3 Roll Ortalaması:", 40),
    "AJ": ("Konum 3 Pitch Ortalaması:", 40),
    "AK": ("Euler Test Sonucu:", 40),
    "AL": ("Gyro Z Testi Yaw Farkı:", 40),
    "AM": ("Gyro Z Test Sonucu:", 40),
    

}

Result_Values = {
    "C": ("Boş Olmalı", 30),
    "D": ("Boş Olmalı", 30),
    "F": ("Eşleşmeli", 45),
    "G": ("Eşleşmeli", 45),
    "H": ("Eşleşmeli", 45),
    "I": ("Eşleşmeli", 45),
    "J": ("Yüklü Olmalı", 30),
    "K": ("Yüklü Olmalı", 30),
    "L": ("Yüklü Olmalı", 30),
    "T": ("9.82", 30),
    "U": ("9.78", 30),
    "V": ("Best veya Noted", 50),
    "AE": ("0.15'ten küçük olmalı:", 40),
    "AF": ("0.15'ten küçük olmalı:", 40),
    "AG": ("89.85 - 90 aralığında olmalı:", 40),
    "AH": ("0.15'ten küçük olmalı:", 40),
    "AI": ("179.85 - 180 aralığında olmalı:", 40),
    "AJ": ("0.15'ten küçük olmalı:", 40),
    "AL": ("0.3'ten küçük olmalı:", 40),

}

def _ensure_headers(sheet):

    sheet.row_dimensions[1].height = 30
    for col, (text, width) in HEADERS.items():
        sheet.column_dimensions[col].width = width
        cell = sheet[f"{col}1"]
        if cell.value != text:
            cell.value = text
            cell.font = Font(size=14, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

def _ensure_res_val(sheet):

    sheet.row_dimensions[2].height = 30
    for col, (text, width) in Result_Values.items():
        sheet.column_dimensions[col].width = width
        cell = sheet[f"{col}2"]
        if cell.value != text:
            cell.value = text
            cell.font = Font(size=14, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')


def find_row_by_sn(sheet, sn):
    row = 3
    last_match = None #yeni eklendi
    while True:
        cell_value = sheet.cell(row=row, column=1).value
        if cell_value is None:
            return last_match
        if cell_value == sn:
            #return row
            last_match = row
        row += 1

def write_result(sheet, sn, column, value, red=False):
    row = find_row_by_sn(sheet, sn)
    if row is None:
        #print(f"Could not write result — {sn} not found in column A")
        return
    cell = sheet.cell(row=row, column=column)
    cell.value = value
    if red:
        cell.font=Font(color="FF0000")

def Sonuc(pn, sn, base_path, results_Data_Select, results_Kalibrasyon_Kontrol, results_Reset, results_Acc_Acilis, results_Acc_Dondurme, results_Euler_Kontrol, results_Gyro_Z):
    folder_path = os.path.join(base_path, "Test Sonuclari")
    os.makedirs(folder_path, exist_ok=True)
    excel_folder_path = os.path.join(folder_path, "TestSonuclari.xlsx")

    if os.path.exists(excel_folder_path):
        wb = openpyxl.load_workbook(excel_folder_path)
        sheet = wb[pn] if pn in wb.sheetnames else wb.create_sheet(pn)
    else:
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = pn

    _ensure_headers(sheet)  
    _ensure_res_val(sheet)

    cell = sheet.cell(row=2, column=1)
    value = "Beklenen Sonuç"
    cell.value = value
    cell.font = Font(size=14, bold=True)
    cell.alignment = Alignment(horizontal='center', vertical='center')

    row = 3
    while True:
        cell_value = sheet.cell(row=row, column=1).value

        if cell_value is None:
            sheet.cell(row=row, column=1).value = sn
            break

        #if cell_value == sn:
            #break

        row += 1

    current_time = datetime.datetime.now()

    write_result(sheet, sn, 2, current_time)

    missing = ", ".join(results_Data_Select["missing_columns"])
    unexpect = ", ".join(results_Data_Select["unexpected_columns"])
    write_result(sheet, sn, 3, missing, red=True) 

    write_result(sheet, sn, 4, unexpect, red=True) 

    if results_Data_Select["is_valid"] == True:
        write_result(sheet, sn, 5, "Geçti") 
    else:
        write_result(sheet, sn, 5, "Kaldı", red=True)

    if results_Kalibrasyon_Kontrol["device_sn_match"] == True:
        write_result(sheet, sn, 6, "SN Eşleşti") 
    else:
        write_result(sheet, sn, 6, "SN Eşleşmedi", red=True) 

    if results_Kalibrasyon_Kontrol["firmware_version_match"] == True:
        write_result(sheet, sn, 7, "FW Eşleşti") 
    else:
        write_result(sheet, sn, 7, "FW Eşleşmedi", red=True)

    if results_Kalibrasyon_Kontrol["device_name_match"] == True:
        write_result(sheet, sn, 8, "DN Eşleşti") 
    else:
        write_result(sheet, sn, 8, "DN Eşleşmedi", red=True)

    if results_Kalibrasyon_Kontrol["hw_match"] == True:
        write_result(sheet, sn, 9, "HW Eşleşti") 
    else:
        write_result(sheet, sn, 9, "HW Eşleşmedi", red=True)

    if results_Kalibrasyon_Kontrol["acc_is_calibrated"] == True:
        write_result(sheet, sn, 10, "Acc Kalib Yüklü") 
    else:
        write_result(sheet, sn, 10, "Acc Kalib Yüklü Değil", red=True)

    if results_Kalibrasyon_Kontrol["gyro_is_calibrated"] == True:
        write_result(sheet, sn, 11, "Gyro Kalib Yüklü") 
    else:
        write_result(sheet, sn, 11, "Gyro Kalib Yüklü Değil", red=True)

    if results_Kalibrasyon_Kontrol["magn_is_calibrated"] == True:
        write_result(sheet, sn, 12, "Magn Kalib Yüklü") 
    else:
        write_result(sheet, sn, 12, "Magn Kalib Yüklü Değil", red=True)

    if results_Kalibrasyon_Kontrol["calib_ctrl_success"] == True:
        write_result(sheet, sn, 13, "Geçti") 
    else:
        write_result(sheet, sn, 13, "Kaldı", red=True)
    
    if results_Reset["hard_reset_test_1"] == True:
        write_result(sheet, sn, 14, "Geçti") 
    else:
        write_result(sheet, sn, 14, "Kaldı", red=True)
    
    if results_Reset["hard_reset_test_2"] == True:
        write_result(sheet, sn, 15, "Geçti") 
    else:
        write_result(sheet, sn, 15, "Kaldı", red=True)
    

    if results_Reset["soft_reset_test_1"] == True:
        write_result(sheet, sn, 16, "Geçti") 
    else:
        write_result(sheet, sn, 16, "Kaldı", red=True)

    if results_Reset["soft_reset_test_2"] == True:
        write_result(sheet, sn, 17, "Geçti") 
    else:
        write_result(sheet, sn, 17, "Kaldı", red=True)

    if results_Reset["reset_result"] == True:
        write_result(sheet, sn, 18, "Geçti") 
    else:
        write_result(sheet, sn, 18, "Kaldı", red=True)

    if results_Acc_Acilis["CalibrationSuccess"] == True:
        write_result(sheet, sn, 19, "Geçti") 
    else:
        write_result(sheet, sn, 19, "Kaldı", red=True)

    write_result(sheet, sn, 20, results_Acc_Acilis["max_acc_norm"]) 
    write_result(sheet, sn, 21, results_Acc_Acilis["min_acc_norm"]) 

    acc_status = ", ".join(results_Acc_Acilis["all_accNormStatus_results"])
    
    write_result(sheet, sn, 22, acc_status) 

    if results_Acc_Acilis["gyro_acilis_success"] == True:
        write_result(sheet, sn, 23, "Geçti") 
    else:
        write_result(sheet, sn, 23, "Kaldı", red=True)

    if results_Acc_Dondurme["test_1"] == True:
        write_result(sheet, sn, 24, "Geçti") 
    else:
        write_result(sheet, sn, 24, "Kaldı", red=True)

    if results_Acc_Dondurme["test_2"] == True:
        write_result(sheet, sn, 25, "Geçti") 
    else:
        write_result(sheet, sn, 25, "Kaldı", red=True)

    if results_Acc_Dondurme["test_3"] == True:
        write_result(sheet, sn, 26, "Geçti") 
    else:
        write_result(sheet, sn, 26, "Kaldı", red=True)

    if results_Acc_Dondurme["test_4"] == True:
        write_result(sheet, sn, 27, "Geçti") 
    else:
        write_result(sheet, sn, 27, "Kaldı", red=True)

    if results_Acc_Dondurme["test_5"] == True:
        write_result(sheet, sn, 28, "Geçti") 
    else:
        write_result(sheet, sn, 28, "Kaldı", red=True)

    if results_Acc_Dondurme["test_6"] == True:
        write_result(sheet, sn, 29, "Geçti") 
    else:
        write_result(sheet, sn, 29, "Kaldı", red=True)

    if results_Acc_Dondurme["AccDondurmeSuccess"] == True:
        write_result(sheet, sn, 30, "Geçti") 
    else:
        write_result(sheet, sn, 30, "Kaldı", red=True)

    write_result(sheet, sn, 31, results_Euler_Kontrol["test_1_roll"]) 
    write_result(sheet, sn, 32, results_Euler_Kontrol["test_1_pitch"]) 

    write_result(sheet, sn, 33, results_Euler_Kontrol["test_2_roll"]) 
    write_result(sheet, sn, 34, results_Euler_Kontrol["test_2_pitch"]) 

    write_result(sheet, sn, 35, results_Euler_Kontrol["test_3_roll"]) 
    write_result(sheet, sn, 36, results_Euler_Kontrol["test_3_pitch"]) 

    if results_Euler_Kontrol["EulerSuccess"] == True:
        write_result(sheet, sn, 37, "Geçti") 
    else:
        write_result(sheet, sn, 37, "Kaldı", red=True)

    write_result(sheet, sn, 38, results_Gyro_Z["YawDiff"]) 

    if results_Gyro_Z["GyroZSucess"] == True:
        write_result(sheet, sn, 39, "Geçti") 
    else:
        write_result(sheet, sn, 39, "Kaldı", red=True)

    wb.save(excel_folder_path)